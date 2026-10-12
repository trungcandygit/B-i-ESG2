"""Build submission/AAMJAF_submission/ for the Asian Academy of Management Journal of Accounting and Finance
(Penerbit Universiti Sains Malaysia, ScholarOne, anonymized review).

Follows the AAMJAF submission guidelines (ejournal.usm.my/aamjaf/submission-guidelines, read through search results
on 12 October 2026, ledger D-61): first page with the title and a 100-150 word abstract followed by at most six
keywords; no author names in the manuscript, author details on a separate cover page; author-date citations;
footnotes avoided; tables numbered consecutively, referred to in the text, each on a separate sheet. Body text and
APA references come from manuscript/manuscript.md; journal-specific sections from manuscript/aamjaf.md.
"""
import os, re, shutil, sys, zipfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_sfr
from build_manuscript import (load_numbers, parse_sections, parse_exhibits, blocks, base_doc, add_runs, set_spacing,
                              body_par, page_break, save_doc, MS, META, OUTR, ROOT)
from build_sfr import words, heading, write_body, table_exhibit, figure_exhibit, front, AUTHOR_TOKENS, PR
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

AAM = os.path.join(ROOT, 'manuscript', 'aamjaf.md')
SUB = os.path.join(ROOT, 'submission', 'AAMJAF_submission')


def plain_text(t):
    """APA text of manuscript.md with figure labels spelled out; no italics (author instruction D-25)."""
    t = re.sub(r'\bFig\. (S?\d)', r'Figure \1', t)
    return t.replace('the Supplemental Appendix', 'the supplementary material').replace('Supplemental Appendix', 'supplementary material')


build_sfr.emerald_text = plain_text   # exhibit notes use the same conventions as the body
build_sfr.sfr_source = lambda src: src   # source line as written in manuscript.md (single author)
AUTHOR_TOKENS = AUTHOR_TOKENS + ['Vietnam National University', 'International School', 'Xuan Thuy', '355 347', 'VNU']


def exhibit_notes(d, note, source):
    q = d.add_paragraph(); add_runs(q, 'Note:', bold=True, size=9); q.add_run(' ')
    add_runs(q, plain_text(note), size=9); set_spacing(q, 1.0, 2)
    q = d.add_paragraph(); add_runs(q, 'Source:', bold=True, size=9); q.add_run(' ')
    add_runs(q, source.replace('source:=', ''), size=9); set_spacing(q, 1.0, 2)


build_sfr.exhibit_notes = exhibit_notes


def abstract(d, aam, jel):
    heading(d, 'Abstract', 1)
    p = d.add_paragraph(); add_runs(p, aam['abstract']); set_spacing(p, 2.0, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # GFA item 3: the first page shows only the title, the abstract, and at most six keywords (no JEL codes)
    p = d.add_paragraph(); add_runs(p, 'Keywords:', bold=True); p.add_run(' ' + aam['keywords'].replace('; ', ', '))
    set_spacing(p, 2.0, 4)


def manuscript(out, secs, aam, labels, exhibits, meta=None):
    d = base_doc(); front(d, secs, aam)
    if meta:
        for line in blocks(meta['authors'])[:4]:
            q = d.add_paragraph(); add_runs(q, line); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(q, 1.0, 4)
    abstract(d, aam, secs['jel']); page_break(d)
    body = plain_text(secs['body'].replace('[[COMPANION]]', ''))
    assert '[[' not in body
    write_body(d, body, labels)
    ack = blocks(meta_all['statements'])   # GFA item 4: acknowledgements at the end of the manuscript
    i = ack.index('## Acknowledgements')
    heading(d, 'Acknowledgements', 1); body_par(d, ack[i + 1], indent=False)
    heading(d, 'References', 1)
    for b in blocks(secs['references']):
        p = d.add_paragraph(); add_runs(p, b.replace('*', '')); set_spacing(p, 2.0, 4)
        p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3)
    for ex in exhibits:   # each table and figure on a separate sheet after the references
        page_break(d)
        (table_exhibit if ex['kind'] == 'table' else figure_exhibit)(d, **ex['args'])
    for p in d.paragraphs:   # double spacing for running text
        if p.paragraph_format.line_spacing == 1.5:
            p.paragraph_format.line_spacing = 2.0
    save_doc(d, out)
    return body


def main():
    global meta_all
    nums = load_numbers(); secs = parse_sections(MS, nums); meta = parse_sections(META, nums); meta_all = meta
    aam = parse_sections(AAM, nums)
    aam['cover_letter'] = aam['cover_letter'].replace('[[TITLE]]', secs['title'])
    exhibits = parse_exhibits(secs); exhibits_ia = parse_exhibits(secs, 'exhibits_ia')
    labels = [e['args']['caption'].split('|')[0].strip().replace('Fig. ', 'Figure ') for e in exhibits]
    if os.path.isdir(SUB):
        shutil.rmtree(SUB)
    os.makedirs(SUB)

    body = manuscript(os.path.join(SUB, '02_Manuscript_Anonymized.docx'), secs, aam, labels, exhibits)
    manuscript(os.path.join(SUB, '07_Manuscript_with_Author_Details.docx'), secs, aam, labels, exhibits, meta)

    # 01 Cover page: title, authors, affiliations, declarations
    d = base_doc(); front(d, secs, aam)
    for line in blocks(meta['authors']):
        q = d.add_paragraph(); add_runs(q, line); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(q, 1.0, 4)
    for b in blocks(meta['statements']):
        heading(d, b[3:], 1) if b.startswith('## ') else body_par(d, b, indent=False)
    heading(d, 'Declaration of generative AI use', 1); body_par(d, aam['ai_declaration'], indent=False)
    save_doc(d, os.path.join(SUB, '01_Cover_Page.docx'))

    # 04 Figures as EPS and PNG files
    fd = os.path.join(SUB, '04_Figures'); os.makedirs(fd)
    for ex in [e for e in exhibits if e['kind'] == 'figure']:
        n = ex['args']['caption'].split('|')[0].strip().replace('Fig. ', '')
        for ext in ('.eps', '.png'):
            shutil.copy(os.path.join(OUTR, 'figures', ex['args']['png'].replace('.png', ext)), os.path.join(fd, f'Figure_{n}{ext}'))

    # 05 Supplementary material, 06 cover letter
    d = base_doc()
    p = d.add_paragraph(); add_runs(p, 'Supplementary material', bold=True, size=14); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(); add_runs(p, secs['title'], bold=True); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 1.5, 12)
    body_par(d, plain_text(secs['ia_intro']), indent=False)
    if 'appendix_b' in aam:   # robustness text moved out of the main text to meet the 7,500-word limit
        for b in blocks(aam['appendix_b']):
            if b.startswith('# '):
                heading(d, b[2:], 1)
            elif b.startswith('## '):
                heading(d, b[3:], 2)
            else:
                body_par(d, plain_text(b))
    page_break(d)
    for i, ex in enumerate(exhibits_ia):
        (table_exhibit if ex['kind'] == 'table' else figure_exhibit)(d, **ex['args'])
        if i < len(exhibits_ia) - 1:
            page_break(d)
    save_doc(d, os.path.join(SUB, '05_Supplementary_material.docx'))

    d = base_doc()
    for line in ['Nguyen Van Trung (corresponding author)', 'Academy of Policy and Development, Hanoi, Vietnam',
                 'kontrungcany@gmail.com', '', 'The Editor', 'Asian Academy of Management Journal of Accounting and Finance', '']:
        q = d.add_paragraph(); add_runs(q, line); set_spacing(q, 1.0, 0)
    for para in blocks(aam['cover_letter']):
        body_par(d, para, indent=False)
    save_doc(d, os.path.join(SUB, '06_Cover_Letter.docx'))

    # 08 Replication package (code, aggregate outputs, pre-analysis plan; licensed data excluded)
    rp = os.path.join(SUB, '08_Replication_Package'); os.makedirs(rp)
    for sub in ('R', 'validation'):
        shutil.copytree(os.path.join(PR, sub), os.path.join(rp, sub))
    shutil.copy(os.path.join(PR, 'run_all.R'), rp)
    shutil.copytree(OUTR, os.path.join(rp, 'outputs'), ignore=shutil.ignore_patterns('overlap'))
    shutil.copy(os.path.join(ROOT, 'notes', '02_analysis_plan.md'), os.path.join(rp, 'pre_analysis_plan.md'))
    shutil.copy(os.path.join(ROOT, 'manuscript', 'replication_README.md'), os.path.join(rp, 'README.md'))
    shutil.make_archive(rp, 'zip', rp)
    shutil.copy(os.path.join(ROOT, 'manuscript', 'aamjaf_checklist.md'), os.path.join(SUB, '00_CHECKLIST_NopBai.md'))

    # Checks
    res = []
    abw = words(aam['abstract'])
    res.append(('abstract 100-150 words', 100 <= abw <= 150, abw))
    kw = [k.strip() for k in aam['keywords'].split(';')]
    res.append(('keywords <= 6, alphabetical', len(kw) <= 6 and kw == sorted(kw, key=str.lower), kw))
    res.append(('no "Fig." labels', 'Fig. ' not in body, ''))
    res.append(('no em dash', '—' not in body + aam['abstract'] + aam['cover_letter'], ''))
    refs = blocks(secs['references'])
    firsts = [re.match(r"([^,(]+)", r).group(1).strip() for r in refs]
    miss = [f for f in firsts if f.split()[0].rstrip(".") not in body]
    res.append(('every reference cited in text', not miss, miss))
    z = zipfile.ZipFile(os.path.join(SUB, '02_Manuscript_Anonymized.docx'))
    xml = ''.join(z.read(n).decode('utf-8', 'ignore') for n in z.namelist() if n.endswith('.xml'))
    hits = [t for t in AUTHOR_TOKENS if t in xml]
    res.append(('anonymized manuscript: no author identifiers', not hits, hits))
    # GFA item 4 places acknowledgements at the end of the manuscript; their text carries no identifier (token check above)
    res.append(('anonymized manuscript: no people.xml', 'word/people.xml' not in z.namelist(), ''))
    res.append(('no footnotes', 'word/footnotes.xml' not in z.namelist() or '<w:footnote ' not in z.read('word/footnotes.xml').decode('utf-8', 'ignore').replace('w:type="separator"', '').replace('w:type="continuationSeparator"', ''), ''))
    word_like = len(aam['abstract'].split()) + len(re.sub(r'(?m)^(\$\$.*|#.*)$', '', body).split())
    res.append(('GFA item 8: abstract + main text within 7,500 words (whitespace count, as in Word)', word_like <= 7500, word_like))
    for n, ok, det in res:
        print(('PASS' if ok else 'FAIL'), n, det)
    bw = words(re.sub(r'^\$\$.*$', '', body, flags=re.M))
    print(f'body {bw} words; abstract {abw} words')
    return res


if __name__ == '__main__':
    res = main()
    sys.exit(0 if all(r[1] for r in res) else 1)
