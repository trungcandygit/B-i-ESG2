"""Chuyển đề cương dạng .md (mỗi mục một dòng) sang .docx: Times New Roman 13, cách dòng 1,3, mục con thụt theo cấp."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

src, out = sys.argv[1], sys.argv[2]
lines = [l.strip() for l in open(src, encoding='utf-8') if l.strip()]
words = sorted({w for l in lines for w in re.findall(r"\b[A-Za-z]{3,}\b", l.replace('**', ''))})
print('Từ chỉ gồm chữ không dấu (kiểm tra tiếng Anh):', words)
d = Document()
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
for s in d.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = Cm(3); s.right_margin = Cm(2)
for l in lines:
    bold = center = False
    if l.startswith('#'):
        l = l.lstrip('#').strip(); bold = center = True
    if l.startswith('**') and l.endswith('**') and l.count('**') == 2:
        l = l[2:-2]; bold = True
    p = d.add_paragraph()
    for i, part in enumerate(l.split('**')):
        if part:
            r = p.add_run(part); r.bold = bold or (i % 2 == 1)
    pf = p.paragraph_format; pf.space_after = Pt(3); pf.line_spacing = 1.3
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = re.match(r'^(\d+(?:\.\d+)*)\.\s', l)
    if m and not bold:
        pf.left_indent = Cm(0.75 * m.group(1).count('.'))
    elif l.startswith(('A. ', 'B. ', 'C. ', 'D. ', 'Phụ lục ', 'Tiểu kết')):
        pf.left_indent = Cm(0.75)
d.save(out)
